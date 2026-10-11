"""Build the final train70 report by filling the retained contest form."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import tempfile
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
from lxml import etree as E
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / 'reports/final35_document_work'
REFERENCE = ROOT / 'reports/modeling_document_work/reference.docx'
OUT = ROOT / 'reports/모델링결과서_김형수없는김형수_최종제출.docx'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
NS = {'w': W, 'r': R}


def node(tag, **attrs):
    x = E.Element('{' + W + '}' + tag)
    for k, v in attrs.items():
        x.set('{' + W + '}' + k, str(v))
    return x


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class Form:
    def __init__(self):
        assert sha(REFERENCE) == 'a9f66f14f8546b37d6a3a4c2c7406c328d95bb6f7b56b501abce9e4bfba0b534'
        with ZipFile(REFERENCE) as z:
            self.parts = {n: z.read(n) for n in z.namelist()}
            self.infos = z.infolist()
        self.xml = E.fromstring(self.parts['word/document.xml'])
        self.section = E.tostring(self.xml.find('.//w:sectPr', NS))
        self.rows = self.xml.find('.//w:tbl', NS).findall('w:tr', NS)
        self.body = self.rows[9].find('w:tc', NS)
        p = self.body.findall('w:p', NS)
        self.samples = {'body': p[7], 'heading': p[3], 'sub': p[5]}
        self.rels = E.fromstring(self.parts['word/_rels/document.xml.rels'])
        self.image_count = 0
        self.texts = []
        self.images = []
        self.page_starts = [0]
        self.cover_mode = False
        self.clear(self.body)
        rp = self.rows[9].find('w:trPr', NS)
        if rp is not None:
            for x in list(rp):
                if x.tag in ['{' + W + '}trHeight', '{' + W + '}cantSplit']:
                    rp.remove(x)
        v = self.body.find('w:tcPr/w:vAlign', NS)
        if v is not None:
            v.set('{' + W + '}val', 'top')

    def clear(self, cell):
        for x in list(cell):
            if x.tag != '{' + W + '}tcPr':
                cell.remove(x)

    def paragraph(self, text, role='body', size=None):
        p = deepcopy(self.samples.get(role, self.samples['body']))
        for x in list(p):
            if x.tag != '{' + W + '}pPr':
                p.remove(x)
        pp = p.find('w:pPr', NS)
        if pp is None:
            pp = node('pPr'); p.insert(0, pp)
        for name in ['ind', 'spacing', 'keepNext', 'keepLines', 'widowControl', 'jc']:
            for x in pp.findall('w:' + name, NS):
                pp.remove(x)
        pp.append(node('spacing', before=120 if role in ['heading', 'sub'] else 0,
                       after=80, line=300 if self.cover_mode else 276, lineRule='auto'))
        if not self.cover_mode:
            pp.append(node('snapToGrid', val=0))
        pp.append(node('keepNext', val=0))
        pp.append(node('jc', val='left'))
        r = node('r'); pr = node('rPr')
        pr.append(node('rFonts', ascii='돋움', eastAsia='돋움', hAnsi='돋움'))
        if role in ['heading', 'sub']:
            pr.append(node('b'))
        pr.append(node('color', val='000000'))
        pr.append(node('sz', val=size or (24 if role == 'heading' else (20 if role == 'caption' else 22))))
        r.append(pr); t = node('t'); t.text = text
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r.append(t); p.append(r)
        return p

    def add(self, text, role='body'):
        self.body.append(self.paragraph(text, role))
        self.texts.append(text)

    def table(self, headers, records, widths):
        t = node('tbl'); pr = node('tblPr')
        pr.append(node('tblW', w=sum(widths), type='dxa'))
        borders = node('tblBorders')
        for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            borders.append(node(side, val='single', sz=4, color='D9D9D9'))
        pr.append(borders); pr.append(node('tblLayout', type='fixed')); margins = node('tblCellMar')
        for side in ['top', 'left', 'bottom', 'right']:
            margins.append(node(side, w=75, type='dxa'))
        pr.append(margins); t.append(pr); grid = node('tblGrid')
        for width in widths:
            grid.append(node('gridCol', w=width))
        t.append(grid)
        for index, values in enumerate([headers] + list(records)):
            tr = node('tr'); trp = node('trPr'); trp.append(node('cantSplit'))
            if index == 0:
                trp.append(node('tblHeader'))
            tr.append(trp)
            for col, (value, width) in enumerate(zip(values, widths)):
                tc = node('tc'); tcp = node('tcPr')
                tcp.append(node('tcW', w=width, type='dxa'))
                if index == 0:
                    tcp.append(node('shd', fill='EFEFEF'))
                tcp.append(node('vAlign', val='center'))
                tc.append(tcp); p = self.paragraph(str(value), size=20)
                pp = p.find('w:pPr', NS); pp.find('w:spacing', NS).set('{' + W + '}after', '25')
                if col > 0:
                    pp.find('w:jc', NS).set('{' + W + '}val', 'center')
                if index == 0:
                    p.find('w:r/w:rPr', NS).insert(1, node('b'))
                    p.find('w:r/w:rPr', NS).append(node('shd', val='clear', fill='EFEFEF'))
                tc.append(p); tr.append(tc)
            t.append(tr)
        self.body.append(t); self.add('')
        self.texts.extend([' | '.join(map(str, r)) for r in [headers] + list(records)])

    def figure(self, path, caption, width=5.55):
        path = ROOT / path
        assert path.exists(), path
        self.image_count += 1
        with tempfile.TemporaryDirectory() as temp:
            opaque = Path(temp) / 'figure.png'
            # Matplotlib exports an unnecessary all-opaque alpha channel.
            # Remove it for stable Word compatibility-mode PDF conversion.
            im = Image.open(path)
            if im.mode == 'RGBA':
                assert im.getchannel('A').getextrema() == (255, 255)
            im.convert('RGB').save(opaque, dpi=(300, 300))
            doc = Document(); doc.add_paragraph().add_run().add_picture(str(opaque), width=Inches(width))
            fragment = Path(temp) / 'fragment.docx'; doc.save(fragment)
            with ZipFile(fragment) as z:
                x = E.fromstring(z.read('word/document.xml'))
                p = x.find('w:body/w:p', NS)
                image_name = next(n for n in z.namelist() if n.startswith('word/media/'))
                media = f'word/media/final35_{self.image_count}.png'
                self.parts[media] = z.read(image_name)
        used = [int(x.get('Id')[3:]) for x in self.rels if x.get('Id', '').startswith('rId') and x.get('Id', '')[3:].isdigit()]
        rid = f'rId{max(used, default=0) + 1}'
        p.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip').set('{' + R + '}embed', rid)
        props = p.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr')
        props.set('id', str(200 + self.image_count)); props.set('descr', caption)
        rel = E.SubElement(self.rels, '{' + REL + '}Relationship', Id=rid, Type=R + '/image', Target=media.removeprefix('word/'))
        pp = node('pPr'); pp.append(node('keepNext', val=0)); pp.append(node('snapToGrid', val=0)); p.insert(0, pp)
        self.body.append(p); self.add(caption, 'caption')
        self.images.append({'path': str(path.relative_to(ROOT)), 'sha256': sha(path), 'caption': caption})

    def cover(self):
        self.cover_mode = True
        for index, text in [(2, ''), (3, '김형수없는김형수'), (4, '없음'),
                            (5, 'PyTorch, segmentation_models_pytorch, timm, NumPy, OpenCV, Pillow')]:
            cell = self.rows[index].findall('w:tc', NS)[-1]; self.clear(cell); cell.append(self.paragraph(text))
        cell = self.rows[7].find('w:tc', NS); self.clear(cell)
        lines = ['◦ 아이디어 제안 배경',
                 '고탄성 알루미늄 합금 현미경 이미지의 네 상을 자동 분할한다.',
                 '밝기와 해상도 차이를 정규화하고 서로 다른 두 분할 모델을 결합했다.',
                 '학습 70장으로 학습한 최종 모델은 검증 20장에서 mIoU 80.62%를 기록했다.',
                 '', '', '', '',
                 '◦ 파급효과',
                 '상별 마스크와 면적 분율을 제공해 시편 간 조직 비교와 전문가 검토를 돕는다.',
                 '학습 데이터 내부 5-fold 검증으로 분할에 따른 성능 변동을 함께 확인했다.',
                 '실제 test 10장의 마스크와 재현 코드를 제공해 분석 절차를 연결한다.']
        for text in lines + [''] * 7:
            cell.append(self.paragraph(text))
        self.cover_mode = False

    def split_flow_rows(self):
        # Small consecutive outer rows avoid Word's costly pagination of one
        # multi-page cell containing all nested numeric tables and figures.
        outer = self.rows[9].getparent()
        position = outer.index(self.rows[9])
        blocks = [x for x in self.body if x.tag != '{' + W + '}tcPr']
        groups = []
        pending = []
        for block in blocks:
            pending.append(block)
            keep = block.find('w:pPr/w:keepNext', NS) is not None
            if block.tag == '{' + W + '}tbl':
                keep = True
            if not keep:
                groups.append(pending); pending = []
        if pending:
            groups.append(pending)
        for i, group in enumerate(groups):
            row = deepcopy(self.rows[9]); cell = row.find('w:tc', NS)
            row.find('w:trPr', NS).append(node('cantSplit', val=0))
            self.clear(cell)
            borders = cell.find('w:tcPr/w:tcBorders', NS)
            for side in ['top', 'bottom']:
                edge = borders.find('w:' + side, NS)
                if (side == 'top' and i > 0) or (side == 'bottom' and i < len(groups) - 1):
                    edge.set('{' + W + '}val', 'nil')
            for block in group:
                cell.append(deepcopy(block))
            if cell[-1].tag != '{' + W + '}p':
                cell.append(self.paragraph(''))
            outer.insert(position + i, row)
        outer.remove(self.rows[9])

    def split_pages(self):
        outer = self.rows[9].getparent()
        parent = outer.getparent()
        position = parent.index(outer)
        blocks = [x for x in self.body if x.tag != '{' + W + '}tcPr']
        edges = self.page_starts + [len(blocks)]
        outer.remove(self.rows[8]); outer.remove(self.rows[9])
        cursor = position + 1
        for i, (start, end) in enumerate(zip(edges, edges[1:])):
            br = node('p'); pp = node('pPr')
            pp.append(node('pageBreakBefore'))
            pp.append(node('spacing', before=0, after=0, line=20, lineRule='exact'))
            pp.append(node('snapToGrid', val=0))
            br.append(pp)
            parent.insert(cursor, br); cursor += 1
            segments = []; pending = []
            for block in blocks[start:end]:
                if block.find('.//w:drawing', NS) is not None:
                    if pending:
                        segments.append(pending); pending = []
                    segments.append(block)
                else:
                    pending.append(block)
            if pending:
                segments.append(pending)
            for segment_index, segment in enumerate(segments):
                if isinstance(segment, list):
                    page = node('tbl')
                    for tag in ['tblPr', 'tblGrid']:
                        page.append(deepcopy(outer.find('w:' + tag, NS)))
                    if i == 0 and segment_index == 0:
                        page.append(deepcopy(self.rows[8]))
                    row = deepcopy(self.rows[9]); cell = row.find('w:tc', NS)
                    row.find('w:trPr', NS).append(node('cantSplit', val=0))
                    self.clear(cell)
                    for block in segment:
                        cell.append(deepcopy(block))
                    if cell[-1].tag != '{' + W + '}p':
                        cell.append(self.paragraph(''))
                    page.append(row)
                    parent.insert(cursor, page); cursor += 1
                else:
                    picture = deepcopy(segment)
                    pp = picture.find('w:pPr', NS)
                    pp.append(node('jc', val='center'))
                    pp.append(node('spacing', before=0, after=0))
                    parent.insert(cursor, picture); cursor += 1

    def save(self):
        self.split_pages()
        # The retained paragraph prototypes carry source editing identifiers.
        # Each paragraph in the completed document must have its own identity.
        for paragraph in self.xml.iter():
            if paragraph.tag == '{' + W + '}sectPr':
                continue
            for attr in list(paragraph.attrib):
                if E.QName(attr).localname in ['paraId', 'textId'] or E.QName(attr).localname.startswith('rsid'):
                    del paragraph.attrib[attr]
        assert E.tostring(self.xml.find('.//w:sectPr', NS)) == self.section
        self.parts['word/document.xml'] = E.tostring(self.xml, xml_declaration=True, encoding='UTF-8', standalone=True)
        self.parts['word/_rels/document.xml.rels'] = E.tostring(self.rels, xml_declaration=True, encoding='UTF-8', standalone=True)
        settings = E.fromstring(self.parts['word/settings.xml'])
        for compatibility in settings.findall('.//w:compatSetting', NS):
            if compatibility.get('{' + W + '}name') == 'compatibilityMode':
                compatibility.set('{' + W + '}val', '15')
        self.parts['word/settings.xml'] = E.tostring(settings, xml_declaration=True, encoding='UTF-8', standalone=True)
        ct = E.fromstring(self.parts['[Content_Types].xml'])
        if not any(x.get('Extension') == 'png' for x in ct):
            E.SubElement(ct, '{http://schemas.openxmlformats.org/package/2006/content-types}Default', Extension='png', ContentType='image/png')
            self.parts['[Content_Types].xml'] = E.tostring(ct, xml_declaration=True, encoding='UTF-8', standalone=True)
        with ZipFile(OUT, 'w') as z:
            for info in self.infos:
                z.writestr(info, self.parts.pop(info.filename))
            for name, value in self.parts.items():
                z.writestr(name, value)
        with ZipFile(OUT) as z, ZipFile(REFERENCE) as ref:
            fixed = [n for n in ref.namelist() if n not in ['word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml', 'word/settings.xml']]
            assert all(z.read(n) == ref.read(n) for n in fixed)
        (WORK / 'report_text.txt').write_text('\n\n'.join(self.texts), encoding='utf-8')
        evidence = {'output': str(OUT), 'sha256': sha(OUT), 'reference_sha256': sha(REFERENCE),
                    'preserve_only_parts_match': True, 'preserve_only_parts': fixed, 'section_unchanged': True,
                    'official_validation20_mIoU': 80.62075310543737, 'train70_internal_oof_mIoU': 80.77829756479382,
                    'images': self.images, 'visual_qa_complete': False}
        (WORK / 'build_evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps({'output': str(OUT), 'sha256': sha(OUT), 'figures': self.image_count}, ensure_ascii=False))


def main():
    WORK.mkdir(exist_ok=True, parents=True)
    content = json.loads((WORK / 'content.json').read_text(encoding='utf-8'))
    f = Form(); f.cover()
    for index, block in enumerate(content['blocks']):
        if index in [11, 19, 29, 34, 42, 48, 55]:
            f.page_starts.append(len(f.body) - 1)
        if block['kind'] == 'text':
            f.add(block['text'], block.get('role', 'body'))
        elif block['kind'] == 'table':
            f.table(block['headers'], block['rows'], block['widths'])
        elif block['kind'] == 'figure':
            if 'report_eda' in block['path']:
                f.add('표 1  train 70장의 정답 상 분포', 'caption')
                f.table(['정답 상', '전체 픽셀 중 비율 (%)', '해당 상이 있는 사진 수'],
                        [['Al', '57.28', '70 / 70'], ['Al3Ni', '26.31', '70 / 70'],
                         ['Primary Si', '9.43', '69 / 70'], ['Eutectic Si', '6.98', '70 / 70']],
                        [2800, 2800, 2900])
            else:
                cases = json.loads((ROOT / 'reports/final35_evidence/validation_cases.json').read_text(encoding='utf-8'))
                records = [['mIoU (%)'] + [f'{c["mIoU"]:.2f}' for c in cases]]
                for phase in ['Al', 'Al3Ni', 'Primary Si', 'Eutectic Si']:
                    records.append([phase + ' IoU (%)'] + [f'{c["class_IoU"][phase]:.2f}' if c['class_IoU'][phase] is not None else '해당 상 없음' for c in cases])
                records.append(['오분류 픽셀 (%)'] + [f'{c["error_pixel_percent"]:.2f}' for c in cases])
                f.add('표 2  공식 validation의 최저점 및 중앙 부근 사례', 'caption')
                f.table(['평가 항목'] + [c['name'] for c in cases], records, [2500, 3000, 3000])
                f.add('초정 Si가 없는 첫 사례는 정답과 예측의 합집합이 없어 해당 IoU를 계산하지 않았다. 오분류 픽셀 비율은 클래스별 평균인 mIoU와 다른 지표다.', 'caption')
        else:
            raise ValueError(block)
    f.save()


if __name__ == '__main__':
    main()
