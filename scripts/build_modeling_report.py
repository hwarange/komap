"""Populate the retained form using completed training and fixed inference reviews."""
import argparse
from copy import deepcopy
import csv
import hashlib
import json
import math
from pathlib import Path
import statistics
import tempfile
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
from lxml import etree as E

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / 'reports/modeling_document_work'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
NS = {'w': W, 'r': R}


def element(tag, **attrs):
    node = E.Element('{' + W + '}' + tag)
    for name, value in attrs.items():
        node.set('{' + W + '}' + name, str(value))
    return node


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--draft', action='store_true')
    parser.add_argument('--crop-review', action='store_true', help='Include completed crop and selective-blending pilots')
    crop_status = parser.add_mutually_exclusive_group()
    crop_status.add_argument('--full-crop-pending', action='store_true', help='Describe the authorized remaining-fold validation as pending')
    crop_status.add_argument('--full-crop-complete', action='store_true', help='Include verified completed five-fold crop validation')
    parser.add_argument('--rubric-review', action='store_true', help='Add verified practical outputs and sharpen the competition narrative')
    parser.add_argument('--development-plan', action='store_true', help='Add the proposed eutectic-silicon error analysis and regional reinference plan')
    args = parser.parse_args()
    if args.full_crop_pending and not args.crop_review:
        parser.error('--full-crop-pending requires --crop-review')
    if args.full_crop_complete and (args.draft or not args.crop_review):
        parser.error('--full-crop-complete requires completed --crop-review evidence')
    if args.rubric_review and (args.draft or not args.crop_review):
        parser.error('--rubric-review requires completed --crop-review evidence')
    if args.development_plan and not args.rubric_review:
        parser.error('--development-plan requires --rubric-review')
    request = json.loads((WORK / 'request.json').read_text(encoding='utf-8'))
    results = ROOT / request['final_result']
    result = None
    weights_path = ROOT / 'runs/review17_weights/results.json'
    tiles_path = ROOT / 'runs/review17_tiles/results.json'
    weights = tiles = None
    crop = selective = full_crop = None
    crop_path = ROOT / 'runs/review18_crop/results.json'
    selective_path = ROOT / 'runs/review19_selective/crop/results.json'
    full_crop_path = ROOT / 'runs/review20_crop_full/results.json'
    full_crop_done_path = full_crop_path.parent / 'all.done'
    full_crop_rows_path = full_crop_path.parent / 'per_image_rows.json'
    if not args.draft:
        done = json.loads((ROOT / 'runs/review16_additional/all.done').read_text())
        assert done['outcome'] == 'three_seed0_jobs_full_cv_and_final_report_complete'
        result = json.loads(results.read_text(encoding='utf-8'))
        assert result['checked_folds'] == list(range(5)) and not result['further_training_scheduled']
        weights = json.loads(weights_path.read_text(encoding='utf-8'))
        tiles = json.loads(tiles_path.read_text(encoding='utf-8'))
        tile_done = json.loads((ROOT / 'runs/review17_tiles/all.done').read_text(encoding='utf-8'))
        assert not weights['training_started'] and not weights['submission_changed']
        assert not tiles['training'] and not tiles['submission_changed'] and not tile_done['training']
        assert tiles['baseline']['images'] == 90
        if args.crop_review:
            crop_done = json.loads((crop_path.parent / 'all.done').read_text())
            selective_done = json.loads((selective_path.parent / 'all.done').read_text())
            assert sha(crop_path) == crop_done['results_sha256']
            assert sha(selective_path) == selective_done['results_sha256']
            crop = json.loads(crop_path.read_text())
            selective = json.loads(selective_path.read_text())
            assert crop['images'] == selective['images'] == 36
            assert crop['folds'] == selective['folds'] == [0, 1]
            assert not crop['submission_changed'] and not selective['submission_changed']
        if args.full_crop_complete:
            full_crop_done = json.loads(full_crop_done_path.read_text(encoding='utf-8'))
            assert full_crop_done['outcome'] == 'matched_crop_five_fold_validation_complete'
            assert sha(full_crop_path) == full_crop_done['results_sha256']
            assert full_crop_done['epochs'] == 60 and full_crop_done['validation_images'] == 90
            assert not full_crop_done['further_training_scheduled'] and not full_crop_done['submission_changed']
            full_crop = json.loads(full_crop_path.read_text(encoding='utf-8'))
            full_crop_rows = json.loads(full_crop_rows_path.read_text(encoding='utf-8'))
            assert full_crop['folds'] == full_crop_rows['folds'] == list(range(5))
            assert full_crop['images'] == 90
            assert not full_crop['further_training_scheduled'] and not full_crop['submission_changed']
            assert full_crop['matched_design']['epochs_both'] == 60
            assert full_crop['matched_design']['seed_both'] == 0
            assert len(full_crop_rows['fold_of']) == 90
            scores = full_crop['scopes']['full90']['scores']
            assert len(scores) == len(full_crop_rows['rows']) == 9
            for recipe, rows_by_name in full_crop_rows['rows'].items():
                assert set(rows_by_name) == set(full_crop_rows['fold_of'])
                means = [statistics.mean(v for v in row['tta_pp'] if v is not None and math.isfinite(v))
                         for row in rows_by_name.values()]
                assert abs(statistics.mean(means) * 100 - scores[recipe]) < 1e-8
    with ZipFile(WORK / 'reference.docx') as z:
        parts = {name: z.read(name) for name in z.namelist()}
        infos = z.infolist()
    xml = E.fromstring(parts['word/document.xml'])
    original_section = E.tostring(xml.find('.//w:sectPr', NS))
    table = xml.find('.//w:tbl', NS)
    rows = table.findall('w:tr', NS)
    body = rows[9].find('w:tc', NS)
    templates = body.findall('w:p', NS)
    regular, heading, subsection = templates[7], templates[3], templates[5]

    def paragraph(text, role='body', size=None):
        sample = heading if role == 'heading' else subsection if role == 'sub' else regular
        node = deepcopy(sample)
        for child in list(node):
            if child.tag != '{' + W + '}pPr':
                node.remove(child)
        pp = node.find('w:pPr', NS)
        if pp is None:
            pp = element('pPr'); node.insert(0, pp)
        for tag in ['ind', 'spacing', 'keepNext', 'keepLines', 'widowControl', 'jc']:
            for child in pp.findall('w:' + tag, NS):
                pp.remove(child)
        pp.append(element('spacing', before=120 if role != 'body' else 0,
                          after=80, line=312, lineRule='auto'))
        pp.append(element('jc', val='left'))
        if text in ['③ 모델링', '3-1. 배경/알고리즘 소개']:
            pp.append(element('keepNext'))
        # Match the source's flowing table cell; Word can otherwise defer an
        # entire long body row when keep-with-next chains include nested tables.
        run = element('r'); rp = element('rPr')
        rp.append(element('rFonts', ascii='돋움', eastAsia='돋움', hAnsi='돋움'))
        rp.append(element('sz', val=size or (24 if role == 'heading' else 22)))
        rp.append(element('color', val='000000'))
        if role in ['heading', 'sub']:
            rp.append(element('b'))
        run.append(rp); text_node = element('t'); text_node.text = text
        text_node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        run.append(text_node); node.append(run)
        return node

    def set_cell(row, text):
        cell = rows[row].findall('w:tc', NS)[-1]
        for child in list(cell):
            if child.tag != '{' + W + '}tcPr':
                cell.remove(child)
        cell.append(paragraph(text))

    set_cell(2, '')
    set_cell(3, request['team_name'])
    set_cell(4, request['awards'])
    set_cell(5, 'PyTorch, segmentation_models_pytorch, timm, OpenCV, NumPy, scikit-learn, Pillow')
    summary = rows[7].find('w:tc', NS)
    for child in list(summary):
        if child.tag != '{' + W + '}tcPr':
            summary.remove(child)
    for text in ['◦ 아이디어 제안 배경',
                 '광학현미경 이미지의 Al, Al3Ni, Primary Si, Eutectic Si 자동 분할을 수행했다.',
                 '얇은 조직의 두께와 촬영 대비를 분석해 입력 배율과 학습 조각 크기를 설계했다.',
                 '합금 단위 교차검증을 거쳐 2배 MiT와 3배 UNet++를 결합한 sub6을 채택했다.']:
        summary.append(paragraph(text))
    for _ in range(4):
        summary.append(paragraph(''))
    for text in ['◦ 파급효과',
                 '예측 마스크를 상별 면적 분율로 요약해 현미경 조직 분석의 검토 자료로 제공한다.',
                 'test 10장의 분율 산출과 도표 생성을 구현해 결과의 활용 절차를 제시했다.' if args.rubric_review else '촬영 조건이 다른 사진의 정량 비교와 전문가 검토를 지원하는 활용안을 제시한다.',
                 '예측 오차를 함께 확인하며 소재의 조성이나 탄성률을 직접 측정한 값과 구분한다.']:
        summary.append(paragraph(text))
    for _ in range(7):
        summary.append(paragraph(''))
    for child in list(body):
        if child.tag != '{' + W + '}tcPr':
            body.remove(child)
    # The empty form's minimum-height body row must expand to the actual report.
    row_properties = rows[9].find('w:trPr', NS)
    if row_properties is not None:
        for tag in ['trHeight', 'cantSplit']:
            for node in row_properties.findall('w:' + tag, NS):
                row_properties.remove(node)
    vertical = body.find('w:tcPr/w:vAlign', NS)
    if vertical is not None:
        vertical.set('{' + W + '}val', 'top')

    def add(text, role='body'):
        body.append(paragraph(text, role))

    def numeric_table(headers, data, widths):
        t = element('tbl'); pr = element('tblPr')
        pr.append(element('tblW', w=sum(widths), type='dxa'))
        borders = element('tblBorders')
        for tag in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            borders.append(element(tag, val='single', sz=4, color='D9D9D9'))
        pr.append(borders); margins = element('tblCellMar')
        for tag in ['top', 'bottom', 'left', 'right']:
            margins.append(element(tag, w=80, type='dxa'))
        pr.append(margins); t.append(pr)
        grid = element('tblGrid')
        for width in widths:
            grid.append(element('gridCol', w=width))
        t.append(grid)
        for index, values in enumerate([headers, *data]):
            tr = element('tr')
            if index == 0:
                tp = element('trPr'); tp.append(element('tblHeader')); tr.append(tp)
            for col, (value, width) in enumerate(zip(values, widths)):
                tc = element('tc'); tcp = element('tcPr')
                tcp.append(element('tcW', w=width, type='dxa'))
                tcp.append(element('vAlign', val='center'))
                if index == 0:
                    tcp.append(element('shd', fill='EFEFEF'))
                tc.append(tcp)
                para = paragraph(str(value), size=20)
                pp = para.find('w:pPr', NS)
                pp.find('w:spacing', NS).set('{' + W + '}after', '30')
                if col > 0:
                    pp.find('w:jc', NS).set('{' + W + '}val', 'center')
                if index == 0:
                    para.find('w:r/w:rPr', NS).append(element('b'))
                tc.append(para); tr.append(tc)
            t.append(tr)
        body.append(t); add('')

    rels = E.fromstring(parts['word/_rels/document.xml.rels'])
    image_number = 0

    def figure(path, caption):
        nonlocal image_number
        image_number += 1
        fragment = Document()
        fragment.add_paragraph().add_run().add_picture(str(ROOT / path), width=Inches(5.6))
        with tempfile.TemporaryDirectory() as temp:
            frag_path = Path(temp) / 'fragment.docx'; fragment.save(frag_path)
            with ZipFile(frag_path) as z:
                root = E.fromstring(z.read('word/document.xml'))
                para = root.find('w:body/w:p', NS)
                image_name = next(n for n in z.namelist() if n.startswith('word/media/'))
                media_name = f'word/media/modeling_report_{image_number}.png'
                parts[media_name] = z.read(image_name)
        relid = f'rIdReportImage{image_number}'
        blip = para.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
        blip.set('{' + R + '}embed', relid)
        props = para.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr')
        props.set('id', str(100 + image_number)); props.set('descr', caption)
        rel = E.SubElement(rels, '{' + REL + '}Relationship')
        rel.set('Id', relid); rel.set('Type', R + '/image')
        rel.set('Target', f'media/modeling_report_{image_number}.png')
        image_pp = para.find('w:pPr', NS)
        if image_pp is None:
            image_pp = element('pPr'); para.insert(0, image_pp)
        image_pp.append(element('keepNext'))
        body.append(para); add(caption, 'caption')

    add('① 분석(EDA) 개요', 'heading')
    add('1-1. 분석 핵심 내용', 'sub')
    add('본 과제는 고탄성 알루미늄 합금의 현미경 이미지를 네 상으로 분할하는 것이다. 촬영 대비·해상도 차이와 얇은 Eutectic Si의 경계 손실에 대응해, 밝기 정규화·입력 배율·모델 조합을 비교 실험으로 검증했다.')
    if tiles:
        highest = (full_crop['scopes']['full90']['scores']['crop1536_uniform'] if full_crop
                   else tiles['tile1536_addition_vs_sub6']['full']['cv_mIoU'])
        add(f"현재 채택 모델은 2배 MiT와 3배 UNet++의 동일 비중 앙상블 sub6이며, 합금 단위 90장 교차검증 mIoU는 {tiles['baseline']['cv_mIoU']:.4f}%다. 후속 4배 모델 조합의 최고 관측값 {highest:.4f}%는 탐색 결과로 구분했다. 평균 점수뿐 아니라 합금별 변화와 추가 추론 비용을 확인해 채택 여부를 판단했다.")
    else:
        add('기준 조합은 2배 MiT와 3배 UNet++의 동일 비중 앙상블인 sub6이다. 4배 MiT seed 0의 전체 5-fold 검증과 고정 비중·추론 타일 비교 결과를 반영한다.')
    add('1-2. 분석 배경', 'sub')
    add('같은 소재의 상은 밝기만으로 구분하기 어렵고 얇은 경계는 작은 픽셀 이동에도 영향을 받는다. 따라서 평균 점수와 함께 상별·저대비 사진의 오차와 합금 간 성능 차이를 확인했다.')
    add('1-3. 분석 방법 및 분석 내용', 'sub')
    numeric_table(['데이터', '사진 수', '정답 마스크', '활용'], [
        ['train', '70', '있음', '초기 학습 및 교차검증'],
        ['validation', '20', '있음', '초기 고정분할 평가 및 교차검증'],
        ['test', '10', '없음', '최종 마스크 생성 및 형식 확인']], [1800, 1000, 1400, 4300])
    add('원본의 밝기 분위수·촬영 스타일·폭과 정답 마스크의 상별 조직 두께를 분석했다. 검증 분할에는 파일명의 합금 번호를 사용했다. 정답이 없는 test는 학습이나 성능 계산에 사용하지 않았다.')
    add('원본 폭은 675·1096·1230·1278·1440픽셀로 나뉜다. EDA에서 675픽셀 사진의 일부 Al3Ni·Eutectic Si 두께가 약 1.9픽셀로 확인되어, 사진 전체를 축소하는 접근보다 배율을 맞추고 타일로 처리하는 접근을 택했다. 기준 모델 오답의 79.1%가 정답 경계 1픽셀 이내에 있었다.')
    if args.rubric_review:
        add('촬영 대비 차이에 맞춰 이미지별 밝기 정규화와 학습 시 밝기 증강을 적용했다. 얇은 조직과 경계 오차에는 확대 입력과 경계 가중 손실로 대응했다. 같은 합금의 사진이 학습과 검증에 함께 들어가지 않도록 합금 번호 단위로 분리했다.')
    figure('eda/figures/fig03_thickness.png', '조직 두께 분석  EDA 원본 그림 3  얇은 구조 보존을 위한 해상도 설계의 근거')
    add('② 데이터 전처리', 'heading')
    add('2-1. 전처리 방법', 'sub')
    add('RGBA로 저장된 흑백 사진을 1채널로 읽고, 각 이미지의 하위 1%와 상위 99% 밝기를 기준으로 정규화했다. 이 정규화는 학습과 추론에 동일하게 적용했다. 촬영 조건을 다양하게 만드는 무작위 밝기 증강은 학습에만 사용했다. 지정된 네 색의 마스크는 클래스 번호 0~3으로 변환하며 밝기 정규화는 적용하지 않았다.')
    add('배율은 기준 폭에 가장 가까운 정수 배율로 정한다. 2배 MiT는 기준 폭 2560, 4배 MiT는 같은 기준 폭에 scale_multiplier=2를 적용한다. 따라서 대부분 사진은 각각 2배·4배이고, 폭 675픽셀 사진은 4배·8배다. 확대는 기존 픽셀의 표현을 바꾸며 원본에 없는 관측 정보를 생성하지 않는다.')
    add('2-2. 전처리 절차', 'sub')
    add('① 이미지와 마스크의 이름·크기·클래스 확인 → ② 흑백 변환 및 밝기 정규화 → ③ 합금 번호 단위 fold 분리 → ④ 배율을 반영한 무작위 조각 추출과 증강 → ⑤ 학습 → ⑥ 타일 예측 및 원본 크기 복원 → ⑦ 후처리와 지표 계산 순서로 처리했다.')
    add('학습에는 무작위 자르기, 크기 증강, 90도 회전·반전 및 밝기 증강을 사용했다. 이미지 보간은 선형, 정답 마스크 보간은 최근접 방식으로 클래스 경계를 처리한다. 이번 추가 검증은 기존 데이터 180개 파일의 SHA256 일치를 확인하고 동일한 fold 파일을 재사용했다.')
    # The source template uses one flowing table cell. Its keep-with-next
    # setting is not honored at this boundary by Word; retain heading space.
    if not args.rubric_review:
        add('')
        add('')
    else:
        add('')
        add('')
        add('')
    add('③ 모델링', 'heading')
    add('3-1. 배경/알고리즘 소개', 'sub')
    add('U-Net 계열은 인코더에서 얻은 특징과 디코더의 위치 정보를 결합해 원본 크기의 분할 지도를 예측한다. MiT-B2 인코더를 사용하는 U-Net과 ResNet34 인코더를 사용하는 UNet++를 비교하고, 서로 다른 오류를 보완하도록 두 모델의 클래스 확률을 평균했다.')
    add('4배 MiT의 원리는 모델의 파라미터를 늘리는 것이 아니라 더 확대된 조직을 보여주는 것이다. 기존 배율 비교에서 입력 조각은 2배와 4배 모두 1024×1024였다. 대부분 사진에서 2배는 원본 약 512×512 영역을, 4배는 약 256×256 영역을 본다. 얇은 조직의 표현은 커지지만 주변 맥락은 줄어들 수 있어 검증이 필요하다.')
    if args.rubric_review:
        add('본 접근의 차별점은 현미경 조직의 두께와 촬영 해상도를 근거로 입력 표현을 설계하고, 서로 다른 모델의 결합 효과를 같은 검증 조건에서 확인한 데 있다. 기존 U-Net·UNet++·MiT 구조를 활용했으며 새 신경망 구조를 개발한 것은 아니다. 후속 crop1536은 확대 배율을 유지하면서 주변 맥락을 넓히는 시도다. crop 면적도 2.25배 늘어나므로 같은 에폭이라도 처리 픽셀 양이 달라지며, 효과를 맥락 하나로 단정하지 않는다.')
    add('3-2. 모델링 방법', 'sub')
    add('초기 전처리와 증강을 검토한 뒤 입력 배율과 모델 구조를 개선하고, 학습 기간 및 앙상블 기여를 확인했다. 주요 결과는 다음과 같다.')
    numeric_table(['주요 실험', '결과', '해석'], [
        ['밝기 증강 제거', '−1.03%p', '해당 대조 실험에서 악화'],
        ['2배 MiT / 60 epoch', '80.171%', '최종 조합의 첫 모델'],
        ['3배 UNet++ / 120 epoch', '80.162%', '최종 조합의 둘째 모델'],
        ['두 모델 동일 비중 sub6', '80.7884%', '채택 조합'],
        ['2배 MiT / 120 epoch', '80.190%', '60 epoch 대비 이득 불확실']], [3700, 1800, 3000])
    add('밝기 증강 제거는 당시 기준 대비 변화이며, 나머지는 합금 단위 90장 CV 점수다. MiT의 120 epoch 연장은 약 +0.019%p에 그쳐 개선 근거가 부족했다.')
    if args.rubric_review:
        add('기존 4배 MiT는 ImageNet으로 초기화한 U-Net / MiT-B2를 seed 0, 60 epoch, batch 4, 사진당 8조각으로 학습했다. crop/tile 1024, width 2560, scale_multiplier 2를 사용했다. AdamW(lr 0.0003, weight decay 0.0001), 교차엔트로피와 Dice 손실, 경계 가중 bw=2를 적용했다. L40S 1개와 BF16으로 학습하고 마지막 epoch를 8방향 TTA 및 overlap 128로 평가했다.')
    else:
        numeric_table(['항목', '기존 4배 MiT 전체 검증 설정'], [
            ['모델과 학습', 'U-Net / MiT-B2, ImageNet 초기화; seed 0, 60 epoch, batch 4, 사진당 8조각'],
            ['입력과 배율', 'crop/tile 1024, width 2560, scale_multiplier 2'],
            ['최적화와 손실', 'AdamW, lr 0.0003, weight decay 0.0001; 교차엔트로피 + Dice, 경계 가중 bw=2'],
            ['평가와 환경', '마지막 60 epoch, 8방향 TTA, overlap 128; L40S 1개, BF16']], [2200, 6300])
    add('합금 번호 단위로 90장을 5개 fold에 나눠, 매번 72장으로 학습하고 해당 모델이 보지 않은 18장으로 평가한다. 같은 합금의 사진은 서로 다른 fold에 나누지 않았다. 앞선 시드 비교에서는 비용을 줄이기 위해 seed 1의 남은 세 fold를 생략하고 seed 0만 전체 검증했다.')
    add('4배 학습 후 주 후보는 2배 MiT·3배 UNet++·4배 MiT seed 0을 각각 1/3로 평균해 비교했다. 기존 sub6은 앞의 두 모델을 각각 1/2로 평균한다. 이후 재학습 없이 고정 비중 두 가지와 추론 타일 1536을 추가 비교했다. 이 후속 검토는 이미 관찰한 90장에 대한 탐색이며 독립 검증으로 간주하지 않았다.')
    add('3-3. 성능', 'sub')
    with (ROOT / 'runs/baseline/analysis/per_image.csv').open(encoding='utf-8-sig') as f:
        validation_rows = list(csv.DictReader(f))
    assert len(validation_rows) == 20
    valid_score = sum(float(row['tta+pp']) for row in validation_rows) / 20
    add('가. 초기 기준 모델의 고정분할 평가', 'sub')
    add(f'초기 U-Net + ResNet34 기준 모델은 train 70장으로 학습하고 validation 20장의 분할 마스크를 예측했다. 이미지별 네 상 IoU를 평균한 후 20장에 대해 평균한 TTA 및 후처리 성능은 {valid_score:.4f}%다. 이 값은 기준 모델의 검증 기록이며 최종 sub6의 고정분할 성능으로 표시하지 않는다. 기준 모델의 체크포인트 선택에 validation을 사용해 낙관적인 추정이 포함될 수 있다.')
    numeric_table(['모델 및 평가', '사진 수', 'mIoU (%)'], [
        ['train70 학습 기준 모델 / validation20', '20', f'{valid_score:.4f}'],
        ['기존 sub6 / 합금 단위 OOF CV', '90', f"{tiles['baseline']['cv_mIoU']:.4f}" if tiles else '80.7884 (기존 기록)']], [5300, 1200, 2000])
    add('mIoU는 각 이미지에서 클래스별 교집합/합집합을 계산해 평균한 뒤 이미지들에 대해 다시 평균한다. 정답과 예측 모두에 없는 클래스는 해당 이미지의 평균에서 제외한다. 고정분할 20장과 교차검증 90장은 평가 대상 및 학습 구성이 달라 두 점수의 차이를 개선 폭으로 해석하지 않는다.')
    add('나. 4배 MiT seed 0 전체 교차검증', 'sub')
    if result:
        comparisons = [('seed0_vs_MiT', '4배 seed0 단일'), ('add_seed0_equal_three', 'sub6에 4배 seed0 추가'),
                       ('replace_MiT_seed0', 'sub6 MiT를 4배로 교체')]
        data = []
        for key, label in comparisons:
            full, held = result[key]['full'], result[key]['nonpilot']
            data.append([label, f"{full['cv_mIoU']:.4f}", f"{full['delta']:+.4f}", f"{held['delta']:+.4f}"])
        numeric_table(['비교', '전체 mIoU (%)', '전체 변화 (%p)', '새 fold 변화 (%p)'], data, [3200, 1700, 1700, 1900])
        add('단일 모델의 변화는 2배 MiT 대비이며, 추가·교체 조합의 변화는 기존 sub6 대비다. 새 fold는 이번에 완성한 fold 2·3·4의 54장이다.')
        candidate = result['add_seed0_equal_three']; full, held = candidate['full'], candidate['nonpilot']
        lo, hi = full['group_bootstrap_ci95']; hlo, hhi = held['group_bootstrap_ci95']
        add(f'추가 조합의 개선 폭 95% 구간은 전체 [{lo:+.4f}, {hi:+.4f}]%p, 새 fold [{hlo:+.4f}, {hhi:+.4f}]%p다. fold 안에서 합금 그룹을 재표집한 bootstrap 10,000회로 계산했다. 평균 개선과 함께 구간에 0이 포함되는지도 판단한다.')
        numeric_table(['상 또는 조건', '전체 변화 (%p)', '새 fold 변화 (%p)'],
            [[name, f'{delta:+.4f}', f"{held['delta_per_class'][name]:+.4f}"]
             for name, delta in full['delta_per_class'].items()] +
            [['저대비', f"{full['delta_low_contrast']:+.4f}", f"{held['delta_low_contrast']:+.4f}"]], [3900, 2300, 2300])
        fold_changes = ', '.join(f"fold {f} {d:+.4f}" for f, d in full['delta_per_fold'].items())
        add('초기 두 fold의 개선이 전체 평균 상승을 이끌었지만, 추가 fold 중 2와 4에서는 소폭 하락했다.' if args.rubric_review else f'추가 조합의 fold별 변화는 {fold_changes}%p였다. 초기 두 fold의 개선이 전체 평균 상승을 이끌었고, 새 fold 중 2와 4에서는 소폭 하락했다.')
    else:
        add('전체 seed 0 검증이 진행 중이다. 전체 및 새 fold의 점수, 상별·저대비 변화와 95% 구간은 학습 및 OOF 평가가 모두 완료된 결과로 반영한다.')
    add('다. 재학습 없는 고정 비중 및 추론 타일 비교', 'sub')
    if weights and tiles:
        a = weights['comparisons']['candidate_A_x4_0_1']
        b = weights['comparisons']['candidate_B_x4_0_2']
        tile = tiles['tile1536_addition_vs_sub6']
        latest_rows = []
        for label, candidate, held_key in [
                ('고정 비중 A  0.4 / 0.1 / 0.5', a, 'nonpilot'),
                ('고정 비중 B  0.3 / 0.2 / 0.5', b, 'nonpilot'),
                ('4배 타일 1536  각 1/3', tile, 'previous_nonpilot')]:
            f, h = candidate['full'], candidate[held_key]
            latest_rows.append([label, f"{f['cv_mIoU']:.4f}", f"{f['delta']:+.4f}", f"{h['delta']:+.4f}"])
        numeric_table(['구성', '전체 mIoU (%)', 'sub6 대비 (%p)', '기존 54장 변화 (%p)'], latest_rows, [3200, 1700, 1700, 1900])
        add('비중은 2배 MiT / 4배 MiT / 3배 UNet++ 순서다. A와 B는 UNet++ 비중을 0.5로 유지한 고정 비교이며 격자 탐색은 하지 않았다. 모든 후보의 전체 개선 95% 구간은 0을 포함했다. 기존 54장은 앞선 실험의 fold 2·3·4로, 이번에는 미관측 검증 데이터가 아니다.')
        tf, th = tile['full'], tile['previous_nonpilot']
        tlo, thi = tf['group_bootstrap_ci95']
        tile_change = tiles['tile1536_addition_vs_tile1024_addition']['full']['delta']
        historical_prefix = '당시 ' if full_crop else ''
        add(f"추론 타일만 1024에서 1536으로 넓힌 추가 조합이 {historical_prefix}최고 {tf['cv_mIoU']:.4f}%를 기록했다. sub6 대비 95% 구간은 [{tlo:+.5f}, {thi:+.5f}]%p, 기존 54장 변화는 {th['delta']:+.5f}%p였다. 같은 추가 조합에서 타일 변경만의 이득은 {tile_change:+.6f}%p로 작았다. 타일 크기는 주변 맥락뿐 아니라 경계 위치와 확률 혼합도 바꾸므로 개선 원인을 하나로 확정하지 않았다.")
        add('같은 표 안의 비교는 동일한 저장 OOF와 평가 코드로 재계산했다.' if args.rubric_review else f"이 표는 원본 서버의 동일 기준 OOF로 재계산했다. sub6은 {tiles['baseline']['cv_mIoU']:.4f}%이며 앞선 RunPod 재계산 {result['baseline']['cv_mIoU']:.4f}%와의 차이는 {tiles['baseline_drift_pp']:+.6f}%p다. 같은 표 안에서 동일 기준과 비교했으며 계산 환경의 미세한 차이를 모델 개선으로 해석하지 않았다.")
    else:
        add('고정 비중 및 추론 타일 비교의 완료 결과를 최종본에 반영한다.')
    if not args.rubric_review:
        add('')
    add('라. 학습 조각과 선택적 혼합의 전체 검증' if full_crop else ('라. 시드와 학습 조각 및 선택적 혼합의 예비 검증' if crop else '라. 예비 결과와 시드 비교 범위'), 'sub')
    if args.rubric_review:
        add('시드 비교는 fold 0·1의 36장에 한정했다. 평균 차이는 작았지만 fold별 상승과 하락이 상쇄됐고 seed 1의 나머지 fold는 수행하지 않았다. 따라서 전체 데이터의 시드 안정성은 아직 확인하지 못했다.' + (' 이후 crop 비교는 seed 0으로 전체 90장 검증을 완료했다.' if full_crop else ' 아래 36장 점수는 전체 90장 점수와 직접 비교하지 않는다.'))
    else:
        add('기존 fold 0·1의 36장에서 4배 seed 0은 81.3048%, seed 1은 81.3063%로 평균 차이가 약 +0.0015%p였다. 그러나 fold별 차이 −0.3219%p와 +0.3248%p가 상쇄된 값이다. seed 1의 나머지 fold는 수행하지 않았으므로 전체 데이터에 대한 시드 안정성을 입증했다고 표현하지 않는다.')
        add('기존 두 시드 평균을 추가한 예비 조합의 +0.3069%p 개선은 이번 seed 0 단독 추가 조합과 구성이 다르다. 예비 점수나 36장의 점수를 이번 전체 90장 결과와 직접 빼서 비교하지 않는다.')
    if full_crop:
        f90 = full_crop['scopes']['full90']
        r54 = full_crop['scopes']['remaining54']
        fc = f90['comparisons']
        single = fc['candidate_vs_control_single']
        uniform = fc['candidate_uniform_vs_control_uniform']
        vs_sub6 = fc['candidate_uniform_vs_sub6']
        numeric_table(['구성', '전체 90장 (%)', '추가 54장 (%)'], [
            [label, f"{f90['scores'][key]:.4f}", f"{r54['scores'][key]:.4f}"]
            for label, key in [('기존 sub6', 'sub6'), ('crop1024 대조군 전체 혼합', 'matched_crop1024_uniform'),
                               ('crop1536 후보 전체 혼합', 'crop1536_uniform'), ('crop1536 선택적 혼합', 'crop1536_selected')]],
            [4100, 2200, 2200])
        slo, shi = single['group_bootstrap_ci95']
        add(f"두 crop 조건 모두 seed 0, 60 epoch, 실제 batch 2·누적 2, 추론 tile1536을 적용하고 마지막 epoch를 평가했다. 예비 36장을 보존하고 fold 2·3·4의 54장을 추가했다. 단일 4배 MiT에서 crop1536은 crop1024보다 {single['delta']:+.4f}%p 높고 합금 그룹 bootstrap 95% 구간은 [{slo:+.4f}, {shi:+.4f}]%p였다. 5개 fold 모두 상승했고 공정 규소는 {single['delta_per_class']['Eutectic Si']:+.4f}%p였다.")
        ulo, uhi = uniform['group_bootstrap_ci95']
        blo, bhi = vs_sub6['group_bootstrap_ci95']
        add(f"세 모델을 각 1/3로 섞으면 후보의 이득은 대조군 대비 {uniform['delta']:+.4f}%p, sub6 대비 {vs_sub6['delta']:+.4f}%p로 줄었다. 각각의 95% 구간 [{ulo:+.4f}, {uhi:+.4f}]%p와 [{blo:+.4f}, {bhi:+.4f}]%p는 0을 포함한다. 추가 54장에서 sub6 대비 {r54['comparisons']['candidate_uniform_vs_sub6']['delta']:+.4f}%p여서 예비 36장의 개선은 그대로 재현되지 않았다.")
        add(f"선택적 혼합은 sub6의 상위 두 클래스 확률 차이 0.15 미만에서만 세 모델 평균을 적용했다. 전체 혼합 대비 {fc['candidate_selected_vs_uniform']['delta']:+.4f}%p로 채택하지 않았다. 저장된 확률의 혼합이며 실제 부분 재추론은 아니다. 반복 검토한 90장의 내부 비교로, 독립 검증이나 시드 안정성의 증거로 해석하지 않는다.")
    elif crop and selective:
        comparisons = crop['comparisons']
        matched = comparisons['crop1536_addition_vs_matched_crop1024_addition']
        historical = comparisons['crop1536_addition_vs_original_tile1536_addition']
        selective_delta = selective['comparisons']['crop1536_selected_vs_uniform']
        numeric_table(['동일한 fold 0·1의 36장 비교', 'mIoU (%)'], [
            ['기존 sub6', f"{crop['sub6_reference']['cv_mIoU']:.4f}"],
            ['새 crop1024 대조군 전체 혼합', f"{matched['base_cv_mIoU']:.4f}"],
            ['새 crop1536 후보 전체 혼합', f"{matched['cv_mIoU']:.4f}"],
            ['crop1536 선택적 혼합', f"{selective_delta['cv_mIoU']:.4f}"]], [6100, 2400])
        lo, hi = matched['group_bootstrap_ci95']
        hlo, hhi = historical['group_bootstrap_ci95']
        add(f"학습 조각은 crop1024와 crop1536으로 비교했다. 두 조건 모두 seed 0, 60 epoch, 실제 batch 2·누적 2, 추론 tile1536을 적용하고 마지막 epoch를 평가했다. crop1536의 전체 혼합은 대조군보다 {matched['delta']:+.4f}%p 높았으며 합금 그룹 bootstrap 95% 구간은 [{lo:+.4f}, {hi:+.4f}]%p였다. 다만 기존 4배 모델의 같은 36장 혼합 {historical['base_cv_mIoU']:.4f}% 대비 이득은 {historical['delta']:+.4f}%p이고 구간 [{hlo:+.4f}, {hhi:+.4f}]%p에 0이 포함됐다.")
        add(f"선택적 혼합은 sub6의 상위 두 클래스 확률 차이가 0.15 미만인 픽셀에만 세 모델의 동일 비중 평균을 적용했다. 전체 혼합보다 {selective_delta['delta']:.4f}%p 낮아 이번 규칙은 채택하지 않았다. 저장된 예측을 섞은 평가이며 관심 영역을 잘라 확대 재추론하지는 않았다. 반복 검토한 36장의 예비 비교이므로 전체 90장 성능이나 독립 검증으로 해석하지 않는다.")
    if (WORK / 'sub6_oof_examples.png').exists():
        figure('reports/modeling_document_work/sub6_oof_examples.png', '그림 2  채택 모델 sub6의 OOF 예측  90장 중 mIoU 중앙값 부근과 최저 점수 사례를 고정 기준으로 선정')
    else:
        figure('runs/baseline/analysis/worst_A71_step2_sample1.png', '그림 2  초기 기준 모델의 validation 오류 사례  최신 후보의 결과와 구분')
    add('마. 채택 판단과 한계', 'sub')
    if result:
        conclusion = f"완료된 전체 90장 검증의 최고 관측 점수는 {highest:.4f}%였지만, 현재 채택 모델은 sub6({tiles['baseline']['cv_mIoU']:.4f}%)이다. 평균 상승만으로 교체하지 않고 합금별 변동과 개선 구간 및 추가 추론 비용을 함께 판단했다."
        if full_crop:
            conclusion += f" 새 최고값은 종전 최고보다 {fc['candidate_uniform_vs_old_uniform']['delta']:+.4f}%p 높지만, 추가 54장과 개선 구간이 일관된 우위를 뒷받침하지 않아 sub6을 유지했다. crop 변경은 주변 맥락과 학습 픽셀 수도 바꾸므로 단일 원인의 효과로 단정하지 않는다."
        elif args.full_crop_pending:
            conclusion += ' crop1536과 동일 조건 crop1024의 남은 fold 2·3·4를 각각 60 epoch로 학습해 전체 5-fold 검증을 진행하며, 완료 후 채택 여부를 다시 판단한다. 전체 결과는 아직 확정하지 않았다.'
        elif crop:
            conclusion += ' 학습 조각 및 선택적 혼합의 예비 비교만으로 채택 모델을 변경하지 않았다.'
        else:
            conclusion += ' 추가 학습 없이 이번 검토를 종료했다.'
        add(conclusion)
        checks = result['add_seed0_equal_three']['predeclared_checks']
        if not args.rubric_review:
            add(f'앞선 4배 학습 후 동일 비중 조합은 사전 판정 기준 {sum(checks.values())}/{len(checks)}개를 충족했다. 전체 +0.05%p 이상, 전체·새 fold의 구간 하한 양수, 새 fold 각각 개선, 저대비 및 Eutectic Si 악화 없음, 다른 상의 하락 0.1%p 이내를 함께 확인했다.')
    else:
        add('최종 채택 판단은 진행 중인 전체 검증 결과로 확정한다. 보고서용 임시 배치에서는 결과와 결론을 미리 확정하지 않는다.')
    add('추가 앙상블은 모델 수와 구조별 비중도 바꾸므로 개선 전체를 확대율만의 효과로 설명할 수 없다. 확대 모델은 처리할 타일과 예측 비용이 늘어난다. 결과는 상별 면적 분율과 형태 정량화에 활용할 수 있으며, 탄성률이나 조성의 직접 측정값과 구분한다.')
    if args.rubric_review:
        add('바. 상별 면적 분율을 이용한 활용', 'sub')
        add('예측 마스크에서 각 상의 픽셀 수를 전체 픽셀 수로 나눈 뒤 100을 곱해 2차원 면적 분율을 산출했다. test 10장의 기존 sub6 마스크를 입력으로 색상 검사, 상별 픽셀 집계, JSON 저장 및 도표 생성을 구현했다. 아래 값은 예측 마스크의 구성비이며 test 정답을 이용한 정확도 검증 결과가 아니다.')
        application_path = ROOT / 'reports/phase_application/phase_area_summary.json'
        application = json.loads(application_path.read_text(encoding='utf-8'))
        assert application['quantity'] == 'predicted_2d_area_fraction' and application['mask_count'] == 10
        assert application['all_masks_exact_four_color_valid']
        assert sha(ROOT / 'scripts/summarize_phase_masks.py') == application['script_sha256']
        assert len(application['images']) == 10
        for row in application['images']:
            assert sha(ROOT / 'submissions/sub6/masks' / row['file']) == row['sha256']
            assert sum(row['class_pixel_counts'].values()) == row['total_pixels']
            assert abs(sum(row['predicted_area_percent'].values()) - 100) < 1e-8
        figure('reports/phase_application/phase_area_fractions.png', '그림 3  sub6의 test 10장 예측 마스크에서 산출한 2차원 상별 면적 분율')
        example = application['images'][0]
        fractions = example['predicted_area_percent']
        add(f"예를 들어 {example['file'].removesuffix('_mask.png')}의 예측 면적 분율은 Al {fractions['Al']:.2f}%, Al3Ni {fractions['Al3Ni']:.2f}%, Primary Si {fractions['Primary Si']:.2f}%, Eutectic Si {fractions['Eutectic Si']:.2f}%다. 모든 사진에서 네 상의 픽셀 수 합이 전체 픽셀 수와 일치하는지 확인했다.")
        add('분석 담당자는 사진별 분율과 원본·마스크를 함께 검토해 비교할 시편을 정할 수 있다. 향후 경계 오차가 큰 사진의 수동 검토와 촬영 조건별 분율 오차 검증을 추가할 수 있다. 실제 수작업 절감 시간과 분율 측정 정확도는 아직 평가하지 않았다. 단면의 예측 면적 분율을 합금 조성, 3차원 부피 분율 또는 탄성률로 환산하지 않는다.')
    add('사. 제출 결과와 재현' if args.rubric_review else '바. 제출 결과와 재현', 'sub')
    add('sub6은 2배 MiT와 3배 UNet++의 각 5-fold 확률 평균을 1:1로 결합한다. 생성한 test 10장의 RGB 마스크는 원본 크기와 지정된 네 색을 검증했다. sub6의 90장 교차검증과 초기 기준 모델의 validation 20장 점수는 구분하며, sub6 자체의 고정분할 20장 성능은 별도 측정하지 않았다.')
    add('학습 설정·코드·체크포인트·fold·결과를 보관한다. Python 3.12, PyTorch 2.14.1+cu126, segmentation_models_pytorch 0.5.0, timm 1.0.30을 사용했다. MiT 초기화는 고정한 smp-hub/mit_b2.imagenet revision을 재사용하고 SHA256을 확인했다.')
    if crop:
        if args.rubric_review:
            add('환경 YAML, 실행 스크립트와 면적 분율 분석 코드를 포함했다. 원본 환경에서 입력·가중치·핵심 버전 검사를 통과했으며, 새 환경 설치와 전체 GPU 재생성 검증은 아직 수행하지 않았다.')
        else:
            add('자체 학습 체크포인트와 결과·코드·재현 기록을 보관한다. 주요 근거는 baseline/analysis, review13~17, seed0_full_results.json, seed0_artifacts.json 및 학습 조각 비교 review18_crop, 선택적 혼합 review19_selective 결과다. 전체 검증은 완료된 fold 0·1을 보존하고 남은 fold 2·3·4만 동일 설정으로 추가하며, 실행 기록은 review20_crop_full에 보관한다.' if args.full_crop_pending else '자체 학습 체크포인트와 결과·코드·재현 기록을 보관한다. 주요 근거는 baseline/analysis, review13~17, seed0_full_results.json, seed0_artifacts.json, review18_crop 및 review19_selective 결과다.')
    else:
        add('추가 학습은 seed 0의 세 fold로 종료했고, 이어서 재학습 없는 비중·타일 비교만 완료했다. 자체 학습 체크포인트와 결과·코드·재현 기록을 보관하며 새로운 실험을 예약하지 않았다. 주요 근거는 baseline/analysis, review13~17 결과, seed0_full_results.json 및 seed0_artifacts.json이다.')
    if args.development_plan:
        oof_path = ROOT / 'runs/review17_tiles/per_image_rows.json'
        oof = json.loads(oof_path.read_text(encoding='utf-8'))['sub6']
        assert len(oof) == 90
        eutectic_iou = [row['tta_pp'][3] for row in oof.values()]
        assert all(value is not None and math.isfinite(value) for value in eutectic_iou)
        eutectic_mean = statistics.mean(eutectic_iou) * 100
        # Source table flow needs local spacing to keep this final heading
        # from being isolated at the foot of the preceding page.
        add('')
        add('')
        add('아. 공정 규소 분할의 후속 개선 계획', 'sub')
        development_intro = ('전체 검증에서 유지하기로 한 sub6을 기준으로 공정 규소(Eutectic Si)의 오류 원인을 분석한다.' if full_crop else '후속 개발은 전체 검증으로 기준 모델을 확정한 뒤 공정 규소(Eutectic Si)의 오류 원인을 분석하는 순서로 진행한다.')
        add(f'{development_intro} 현재 sub6의 90장 OOF에서 공정 규소의 이미지별 IoU 평균은 {eutectic_mean:.2f}%다. 불규칙성과 얇은 두께의 영향을 구분해 확인하며, 아래 내용은 아직 수행하지 않은 개선 계획이다.')
        numeric_table(['오류 유형', '확인할 지표', '검토할 방향'], [
            ['가는 조직 누락', '두께별 재현율·누락 면적', '미세 구조와 원본 해상도'],
            ['다른 상으로 혼동', '혼동행렬·영역별 오분류', '주변 조직과 전체 형태'],
            ['경계 위치·두께 오차', '경계 IoU·두께 편향', '경계 주변 정밀 보정']], [2400, 3100, 3000])
        add('원본·정답·OOF 예측을 함께 검토해 오류를 위 유형으로 나눈다. 경계가 애매한 사례는 별도로 기록하며 평가 정답을 임의로 수정하지 않는다. 우선순위는 오류의 빈도와 면적, 기존 모델로 보정할 가능성에 따라 정한다.')
        add('다음 후보는 실제 부분 재추론이다. 전체 예측에서 불확실하거나 모델 간 판단이 다른 관심 영역을 선정하고 주변 여백을 포함해 기존 학습 모델로 다시 예측한다. 입력 배율·crop·여백·결합 규칙을 고정하고, 원본 좌표로 복원한 확률을 겹침 영역에서 결합해 이음새를 점검한다. 영역 선택에는 정답을 쓰지 않으며, 사후에 선택 면적과 전체 오류 중 포착한 비율을 계산해 확신하면서 틀린 영역의 누락을 확인한다.')
        add('앞선 선택적 혼합은 저장된 확률을 섞은 실험으로 실제 부분 재추론과 구분한다. 기존의 오답 조각 집중 학습·경계 정답 완화·Skeleton 손실은 시험한 설정에서 개선 근거가 부족했으므로, 동일 조건의 반복은 우선하지 않는다.')
        add('작은 탐색에서 가능성을 확인한 구성만 규칙을 고정해 전체 평가로 확장한다. 같은 사진에서 전체 mIoU, 상별 누락·과검출, 면적 분율의 평균 절대오차와 평균 편향(%p), 사진당 총 처리 시간과 재추론 면적을 비교한다. 다른 상의 악화와 합금별 변동도 함께 확인한다. 이미 반복 검토한 90장의 결과는 내부 비교로 해석하며, 독립 일반화 성능은 새로운 합금·시편 자료로 별도 검증한다.')
    assert E.tostring(xml.find('.//w:sectPr', NS)) == original_section
    parts['word/document.xml'] = E.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True)
    parts['word/_rels/document.xml.rels'] = E.tostring(rels, xml_declaration=True, encoding='UTF-8', standalone=True)
    types = E.fromstring(parts['[Content_Types].xml'])
    ct = 'http://schemas.openxmlformats.org/package/2006/content-types'
    if not any(n.get('Extension') == 'png' for n in types):
        node = E.SubElement(types, '{' + ct + '}Default'); node.set('Extension', 'png'); node.set('ContentType', 'image/png')
        parts['[Content_Types].xml'] = E.tostring(types, xml_declaration=True, encoding='UTF-8', standalone=True)
    output = (WORK / 'layout_draft.docx' if args.draft else
              ROOT / ('reports/모델링결과서_김형수없는김형수_심사기준보완.docx' if args.rubric_review else 'reports/모델링결과서_김형수없는김형수.docx'))
    with ZipFile(output, 'w') as z:
        for info in infos:
            z.writestr(info, parts.pop(info.filename))
        for name, content in parts.items():
            z.writestr(name, content)
    with ZipFile(output) as z, ZipFile(WORK / 'reference.docx') as reference:
        unchanged = [name for name in reference.namelist() if name not in
                     ['word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml']]
        assert all(z.read(name) == reference.read(name) for name in unchanged)
    evidence = {'reference_sha256': sha(WORK / 'reference.docx'), 'source_sha256': sha(WORK / 'reference_original.doc'),
                'output': str(output), 'output_sha256': sha(output), 'draft': args.draft,
                'preserve_only_parts_match': True, 'section_unchanged': True, 'validation20_mIoU': valid_score,
                'result_sha256': sha(results) if result else None,
                'weights_result_sha256': sha(weights_path) if weights else None,
                'tiles_result_sha256': sha(tiles_path) if tiles else None,
                'crop_result_sha256': sha(crop_path) if crop else None,
                'selective_result_sha256': sha(selective_path) if selective else None,
                'full_crop_validation_pending': args.full_crop_pending,
                'full_crop_validation_complete': args.full_crop_complete,
                'full_crop_result_sha256': sha(full_crop_path) if full_crop else None,
                'full_crop_completion_sha256': sha(full_crop_done_path) if full_crop else None,
                'full_crop_per_image_sha256': sha(full_crop_rows_path) if full_crop else None,
                'full_crop_90_image_scores_recomputed': bool(full_crop),
                'rubric_revision': args.rubric_review,
                'development_plan_included': args.development_plan,
                'development_plan_executed': False,
                'development_oof_source_sha256': sha(oof_path) if args.development_plan else None,
                'sub6_eutectic_mean_IoU': eutectic_mean if args.development_plan else None,
                'application_result_sha256': sha(application_path) if args.rubric_review else None,
                'application_figure_sha256': sha(ROOT / 'reports/phase_application/phase_area_fractions.png') if args.rubric_review else None,
                'sub6_oof_figure_used': (WORK / 'sub6_oof_examples.png').exists()}
    (WORK / ('draft_evidence.json' if args.draft else 'final_build_evidence.json')).write_text(
        json.dumps(evidence, indent=2, ensure_ascii=False), encoding='utf-8')
    print(json.dumps(evidence, ensure_ascii=False))


if __name__ == '__main__':
    main()
